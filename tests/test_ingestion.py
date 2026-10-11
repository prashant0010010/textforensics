import io
import unittest
from unittest import mock

from core import ingestion
from core.ingestion import (DocumentTooLargeError, IngestionError, ScannedPDFError, from_docx_bytes, from_pasted_text,
                            from_pdf_bytes, from_txt_bytes, from_upload)
from tests import samples


class TestTextInput(unittest.TestCase):
    def test_empty_input(self):
        with self.assertRaises(IngestionError):
            from_pasted_text("")
        with self.assertRaises(IngestionError):
            from_pasted_text("   \n\n  ")
        with self.assertRaises(IngestionError):
            from_txt_bytes(b"")

    def test_utf8_bom_and_crlf(self):
        r = from_txt_bytes("\ufeffHello world.\r\nSecond line.\r\n".encode("utf-8"))
        self.assertNotIn("\r", r.text)
        self.assertFalse(r.text.startswith("\ufeff"))

    def test_cp1252_fallback_warns(self):
        r = from_txt_bytes("caf\xe9 society".encode("cp1252"))
        self.assertIn("café", r.text)
        self.assertTrue(r.warnings)

    def test_utf16(self):
        r = from_txt_bytes("Hello wörld.".encode("utf-16"))
        self.assertEqual(r.encoding, "utf-16")
        self.assertIn("wörld", r.text)

    def test_binary_rejected(self):
        with self.assertRaises(IngestionError):
            from_txt_bytes(b"\x00\x01\x02" * 100)

    def test_unicode_normalised_to_nfc(self):
        r = from_pasted_text("e\u0301cole " * 10)
        self.assertIn("\u00e9cole", r.text)

    def test_hard_limit_is_explicit_not_silent(self):
        with mock.patch.object(ingestion, "HARD_LIMIT_WORDS", 50):
            with self.assertRaises(DocumentTooLargeError):
                from_pasted_text("word " * 100)

    def test_soft_limit_warns_but_keeps_all_text(self):
        with mock.patch.object(ingestion, "SOFT_LIMIT_WORDS", 50):
            r = from_pasted_text("word " * 100)
        self.assertEqual(len(r.text.split()), 100)
        self.assertTrue(any("analysed in full" in w for w in r.warnings))

    def test_unsupported_extension(self):
        with self.assertRaises(IngestionError):
            from_upload("a.exe", b"abc")


class TestDocx(unittest.TestCase):
    def _make(self):
        import docx
        d = docx.Document()
        d.add_heading("Findings", level=1)
        d.add_paragraph("The first paragraph explains the setting. It has two sentences.")
        d.add_paragraph("An item in a list", style="List Bullet")
        t = d.add_table(rows=2, cols=2)
        t.cell(0, 0).text, t.cell(0, 1).text, t.cell(1, 0).text, t.cell(1, 1).text = "a", "b", "c", "d"
        buf = io.BytesIO()
        d.save(buf)
        return buf.getvalue()

    def test_docx_extraction_and_hints(self):
        r = from_docx_bytes(self._make())
        self.assertIn("The first paragraph explains the setting.", r.text)
        kinds = {k for _, _, k in r.hints}
        self.assertTrue({"heading", "bullet", "table"} <= kinds)
        for s, e, k in r.hints:
            self.assertTrue(r.text[s:e].strip())

    def test_corrupt_docx(self):
        with self.assertRaises(IngestionError):
            from_docx_bytes(b"not a docx")


class TestPdf(unittest.TestCase):
    def _text_pdf(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        y = 800
        for i in range(30):
            c.drawString(60, y, f"Line {i} of an ordinary extractable paragraph about harbours and tides.")
            y -= 14
        c.showPage()
        c.save()
        return buf.getvalue()

    def test_pdf_extraction(self):
        r = from_pdf_bytes(self._text_pdf())
        self.assertIn("harbours and tides", r.text)
        self.assertEqual(r.meta["pages"], 1)

    def test_scanned_pdf_is_not_pretended_to_be_analysed(self):
        from reportlab.lib.pagesizes import A4
        from reportlab.pdfgen import canvas
        buf = io.BytesIO()
        c = canvas.Canvas(buf, pagesize=A4)
        c.rect(50, 50, 300, 300, fill=1)       # drawing only, no text layer
        c.showPage()
        c.save()
        with self.assertRaises(ScannedPDFError) as cm:
            from_pdf_bytes(buf.getvalue())
        self.assertIn("NOT analysed", str(cm.exception))

    def test_corrupt_pdf(self):
        with self.assertRaises(IngestionError):
            from_pdf_bytes(b"%PDF-garbage")


if __name__ == "__main__":
    unittest.main()
